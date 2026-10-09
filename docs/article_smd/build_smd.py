"""Build the conference paper (SMD-UTM 2026 template) in two formats from one source:

    python docs/article_smd/build_smd.py

    articol_smd.md        the text (Romanian), with [@key] citations and format-specific blocks
    references_smd.py     the bibliography, already in the template's format
    template_SMD_UTM_2026.docx   the official template: its header, footer, page setup and styles are reused
    smd_template.tex      the LaTeX layout that reproduces the template

Outputs (docs/article_smd/out/):
    articol_SMD_UTM_2026.docx      Word version built on the official template
    articol_SMD_UTM_2026.tex       LaTeX version (+ figures/ and logo_utm.png) -> compile with XeLaTeX
    articol_SMD_UTM_2026.pdf       PDF compiled from the LaTeX version (if xelatex is available)
    latex_overleaf.zip             everything needed to compile the LaTeX version on Overleaf
Needs pandoc >= 3; python-docx.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(HERE, "out")
NAME = "articol_SMD_UTM_2026"
sys.path.insert(0, HERE)
from references_smd import REFS  # noqa: E402

TITLE = ("Descompunerea bayesiană a incertitudinii pentru avertizarea timpurie a devierii conceptuale, "
         "validată prin XAI")
FIGS = ["fig1_example_trace.png", "fig2_signal_auroc.png", "fig7_bayes_vs_freq.png", "fig3_tradeoff.png",
        "fig5_xai.png"]


# ----------------------------------------------------------------------------- citations
def number_citations(md: str):
    """[@a; @b] (markdown) and [CITE:a] (raw LaTeX) -> [1,2] in order of first appearance."""
    pat = re.compile(r"\[@([^\]]+)\]|\[CITE:([^\]]+)\]")
    order = []
    for m in pat.finditer(md):
        keys = [k.strip().lstrip("@") for k in (m.group(1) or m.group(2)).split(";")]
        for k in keys:
            if k not in REFS:
                raise KeyError(f"unknown reference {k}")
            if k not in order:
                order.append(k)
    num = {k: i + 1 for i, k in enumerate(order)}

    def fmt(m):
        ns = sorted(num[k.strip().lstrip("@")] for k in (m.group(1) or m.group(2)).split(";"))
        parts, i = [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            parts.append(f"{ns[i]}–{ns[j]}" if j - i >= 2 else ",".join(map(str, ns[i:j + 1])))
            i = j + 1
        return "[" + ",".join(parts) + "]"
    return pat.sub(fmt, md), order


def md_to_latex_inline(s: str) -> str:
    s = s.replace("\\", r"\textbackslash{}")
    for a, b in [("&", r"\&"), ("%", r"\%"), ("#", r"\#"), ("_", r"\_"), ("$", r"\$")]:
        s = s.replace(a, b)
    s = re.sub(r"(https?://\S+?)(\s|\)|$)", lambda m: r"\url{" + m.group(1).replace(r"\_", "_") + "}" + m.group(2), s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", s)
    s = re.sub(r"\*(.+?)\*", r"\\textit{\1}", s)
    return s


def references_block(order):
    docx = "\n\n".join(f'::: {{custom-style="Referinta"}}\n{i}\\. {REFS[k]}\n:::' for i, k in enumerate(order, 1))
    items = "\n".join(rf"\item {md_to_latex_inline(REFS[k])}" for k in order)
    latex = ("```{=latex}\n{\\raggedright\n\\begin{enumerate}[label=\\arabic*.,leftmargin=1cm,labelsep=0.25cm,"
             "itemsep=0pt,parsep=0pt,topsep=2pt,align=left]\n" + items + "\n\\end{enumerate}\\par}\n```")
    return f"::: {{.only-docx}}\n{docx}\n:::\n\n::: {{.only-latex}}\n{latex}\n:::\n"


DECL_TEXT = ("Subsemnatul(a), în calitate de autor corespondent, declar că lucrarea cu titlul „{t}” este originală, "
             "reprezintă rezultatul activității autorilor și că toate sursele bibliografice utilizate sunt citate "
             "corespunzător. Îmi asum responsabilitatea pentru integritatea lucrării și corectitudinea informațiilor "
             "prezentate în această lucrare în numele tuturor coautorilor.")


def declaration_block():
    t = DECL_TEXT.format(t=TITLE)
    r = '<w:rPr><w:color w:val="222222"/><w:lang w:val="ro-RO" w:eastAsia="en-US"/></w:rPr>'
    line = ('<w:p><w:pPr><w:spacing w:line="276" w:lineRule="auto"/><w:ind w:left="567" w:hanging="567"/>'
            '<w:jc w:val="both"/></w:pPr><w:r>' + r + '<w:t xml:space="preserve">{}</w:t></w:r></w:p>')
    empty = '<w:p><w:pPr><w:spacing w:line="276" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr></w:p>'
    ox = ['<w:p><w:r><w:br w:type="page"/></w:r></w:p>', empty, empty,
          '<w:p><w:pPr><w:spacing w:line="276" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr><w:r>'
          '<w:rPr><w:b/><w:bCs/><w:color w:val="222222"/><w:lang w:val="ro-RO" w:eastAsia="en-US"/></w:rPr>'
          '<w:t>Declarație de originalitate și asumare a responsabilității</w:t></w:r></w:p>', empty,
          '<w:p><w:pPr><w:spacing w:line="360" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr><w:r>' + r +
          f'<w:t xml:space="preserve">{t}</w:t></w:r></w:p>', empty, empty,
          line.format("Numele și prenumele: ______________________"), empty,
          line.format("Semnătura: ______________________"), empty, line.format("Data: ______________________")]
    latex = (r"\newpage\vspace*{2cm}\begin{center}\textbf{Declarație de originalitate și asumare a "
             r"responsabilității}\end{center}\vspace{6pt}{\setlength{\parindent}{0pt}\setstretch{1.5}" + "\n" +
             md_to_latex_inline(t) + r"\par}\vspace{1.5cm}{\setlength{\parindent}{0pt}"
             r"Numele și prenumele: \rule{6cm}{0.4pt}\\[14pt]Semnătura: \rule{6cm}{0.4pt}\\[14pt]"
             r"Data: \rule{6cm}{0.4pt}\par}")
    return ("::: {.only-docx}\n```{=openxml}\n" + "\n".join(ox) + "\n```\n:::\n\n"
            "::: {.only-latex}\n```{=latex}\n" + latex + "\n```\n:::\n")


# ----------------------------------------------------------------------------- Word reference document
STYLES = r"""
<w:style w:type="paragraph" w:customStyle="1" w:styleId="BodyText"><w:name w:val="Body Text"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/><w:ind w:firstLine="708"/><w:jc w:val="both"/></w:pPr><w:rPr><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="FirstParagraph"><w:name w:val="First Paragraph"/><w:basedOn w:val="BodyText"/><w:next w:val="BodyText"/><w:qFormat/></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Compact"><w:name w:val="Compact"/><w:basedOn w:val="BodyText"/><w:qFormat/></w:style>
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="BodyText"/><w:qFormat/>
<w:pPr><w:keepNext/><w:spacing w:before="240" w:after="0"/><w:ind w:firstLine="708"/><w:jc w:val="both"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:bCs/><w:sz w:val="24"/><w:szCs w:val="24"/><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="BodyText"/><w:qFormat/>
<w:pPr><w:keepNext/><w:spacing w:before="120" w:after="0"/><w:ind w:firstLine="708"/><w:jc w:val="both"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:b/><w:bCs/><w:i/><w:iCs/><w:sz w:val="24"/><w:szCs w:val="24"/><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Rezumat"><w:name w:val="Rezumat"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:tabs><w:tab w:val="left" w:pos="2268"/></w:tabs><w:ind w:right="18"/><w:jc w:val="both"/></w:pPr><w:rPr><w:i/><w:iCs/><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="CuvinteCheie"><w:name w:val="CuvinteCheie"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:spacing w:before="240" w:after="0"/><w:jc w:val="both"/></w:pPr><w:rPr><w:i/><w:iCs/><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Ecuatie"><w:name w:val="Ecuatie"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:tabs><w:tab w:val="center" w:pos="4677"/><w:tab w:val="right" w:pos="9354"/></w:tabs><w:spacing w:before="80" w:after="80"/></w:pPr><w:rPr><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Multumiri"><w:name w:val="Multumiri"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:spacing w:before="240" w:after="0"/><w:ind w:firstLine="708"/><w:jc w:val="both"/></w:pPr><w:rPr><w:lang w:val="ro-RO"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Referinta"><w:name w:val="Referinta"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:ind w:left="567" w:hanging="567"/><w:jc w:val="left"/></w:pPr><w:rPr><w:lang w:val="en-US"/></w:rPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="Figure"><w:name w:val="Figure"/><w:basedOn w:val="Normal"/><w:qFormat/>
<w:pPr><w:keepNext/><w:spacing w:before="120" w:after="60"/><w:jc w:val="center"/></w:pPr></w:style>
<w:style w:type="paragraph" w:customStyle="1" w:styleId="CaptionedFigure"><w:name w:val="Captioned Figure"/><w:basedOn w:val="Figure"/><w:qFormat/></w:style>
<w:style w:type="character" w:customStyle="1" w:styleId="Hyperlink"><w:name w:val="Hyperlink"/><w:rPr><w:color w:val="000000"/></w:rPr></w:style>
"""


def make_reference(path):
    """The official template with an empty body (pandoc keeps its header, footer and page setup) plus the
    paragraph styles pandoc needs, each defined to match the template's direct formatting."""
    tpl = os.path.join(HERE, "template_SMD_UTM_2026.docx")
    with zipfile.ZipFile(tpl) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                s = data.decode("utf8")
                for sid in ["BodyText", "FirstParagraph", "Compact", "Heading1", "Heading2", "Figure",
                            "CaptionedFigure", "Hyperlink"]:
                    s = re.sub(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % sid, "", s, flags=re.S)
                data = s.replace("</w:styles>", STYLES + "</w:styles>").encode("utf8")
            if item.filename == "word/document.xml":
                d = data.decode("utf8")
                sect = re.findall(r"<w:sectPr.*?</w:sectPr>", d, re.S)[-1]
                d = re.sub(r"<w:body>.*</w:body>", "<w:body><w:p/>" + sect + "</w:body>", d, flags=re.S)
                data = d.encode("utf8")
            zout.writestr(item, data)


def _borders(tag="tblBorders"):
    b = OxmlElement(f"w:{tag}")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
        b.append(e)
    return b


def postprocess_docx(path):
    d = Document(path)
    for t in d.tables:                          # template: centred, full width, all borders, 10 pt, single spacing
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(old)
        look = tblPr.find(qn("w:tblLook"))          # schema order: ... tblW, jc, ..., tblBorders, ..., tblLook
        anchor = None
        for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption"):
            anchor = tblPr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(_borders())
        else:
            tblPr.append(_borders())
        w = tblPr.find(qn("w:tblW"))
        if w is None:
            w = OxmlElement("w:tblW"); tblPr.append(w)
        w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")
        for r_i, row in enumerate(t.rows):
            for c_i, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.first_line_indent = Pt(0); pf.space_before = Pt(1); pf.space_after = Pt(1)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (c_i == 0 and r_i > 0) else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(10); run.font.name = "Times New Roman"
    for p in d.paragraphs:                      # figures centred, no first-line indent
        if p._p.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
    d.core_properties.title = TITLE
    d.core_properties.author = "Prenumele Numele"
    d.save(path)
    # keep the template's own settings.xml (pandoc reorders an element that strict OOXML validation rejects)
    tpl = os.path.join(HERE, "template_SMD_UTM_2026.docx")
    with zipfile.ZipFile(tpl) as zt:
        settings = zt.read("word/settings.xml")
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                data = settings
            elif item.filename == "word/numbering.xml":  # pandoc writes 4-digit nsid values; schema wants 8 hex digits
                data = re.sub(rb'(<w:nsid w:val=")([0-9A-Fa-f]{1,7})(")', lambda m: m.group(1) + m.group(2).rjust(8, b"0") + m.group(3), data)
            elif item.filename == "word/document.xml":   # pandoc writes m:sepChr after m:endChr; schema wants it before
                data = re.sub(rb"(<m:endChr [^>]*/>)(<m:sepChr [^>]*/>)", rb"\2\1", data)
            zout.writestr(item, data)
    os.replace(tmp, path)


# ----------------------------------------------------------------------------- character count
def count_characters(docx_path):
    """Characters with spaces from the title to the end of the acknowledgments (references and the
    originality declaration excluded), as Word's 'characters (with spaces)' would count them."""
    d = Document(docx_path)
    texts = []
    body = d.element.body
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            t = "".join(x.text or "" for x in el.iter(qn("w:t")))
            if t.strip() == "Referințe":
                break
            texts.append(t)
        elif el.tag == qn("w:tbl"):
            texts.append(" ".join("".join(x.text or "" for x in c.iter(qn("w:t"))) for c in el.iter(qn("w:tc"))))
    math = sum(len("".join(x.text or "" for x in m.iter(qn("m:t")))) for m in body.iter(qn("m:oMath")))
    return sum(len(t) for t in texts) + math


# ----------------------------------------------------------------------------- main
def main():
    os.makedirs(OUT, exist_ok=True)
    md = open(os.path.join(HERE, "articol_smd.md"), encoding="utf8").read()
    md, order = number_citations(md)
    md = md.replace("[[REFERENCES]]", references_block(order)) + "\n" + declaration_block()
    src = os.path.join(OUT, "_combined.md")
    open(src, "w", encoding="utf8").write(md)
    filt = os.path.join(HERE, "only_format.lua")

    # ---- Word, on the official template
    ref = os.path.join(OUT, "_reference.docx")
    make_reference(ref)
    docx = os.path.join(OUT, NAME + ".docx")
    subprocess.run(["pandoc", src, "-f", "markdown", "-o", docx, "--lua-filter", filt, "--reference-doc", ref,
                    "--resource-path", HERE, "-M", "lang=ro-RO"], check=True)
    postprocess_docx(docx)
    n = count_characters(docx)
    print(f"[docx] {docx}\n       characters with spaces (title -> acknowledgments): {n:,} (template limit 12,000)")

    # ---- LaTeX
    texdir = os.path.join(OUT, "latex")
    os.makedirs(os.path.join(texdir, "figures"), exist_ok=True)
    for f in FIGS:
        shutil.copy(os.path.join(ROOT, "figures", "ro", f), os.path.join(texdir, "figures", f))
    shutil.copy(os.path.join(HERE, "logo_utm.png"), texdir)
    tex = os.path.join(texdir, NAME + ".tex")
    subprocess.run(["pandoc", src, "-f", "markdown", "-t", "latex", "-s", "-o", tex, "--lua-filter", filt,
                    "--template", os.path.join(HERE, "smd_template.tex")], check=True)
    t = open(tex, encoding="utf8").read().replace("\\usepackage{enumitem}", "\\usepackage{enumitem}\n\\usepackage{setspace}")
    open(tex, "w", encoding="utf8").write(t)
    zp = os.path.join(OUT, "latex_overleaf.zip")
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for base, _, files in os.walk(texdir):
            for f in files:
                if f.endswith((".tex", ".png")):
                    full = os.path.join(base, f)
                    z.write(full, os.path.relpath(full, texdir))
    print(f"[tex]  {tex}\n       {zp}")
    if shutil.which("xelatex") and "--no-pdf" not in sys.argv:
        for _ in range(2):
            r = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", NAME + ".tex"], cwd=texdir,
                               capture_output=True, text=True)
        if r.returncode != 0:
            print(r.stdout[-3000:])
            raise SystemExit("xelatex failed")
        shutil.copy(os.path.join(texdir, NAME + ".pdf"), os.path.join(OUT, NAME + ".pdf"))
        print(f"[pdf]  {os.path.join(OUT, NAME + '.pdf')}")
    for f in (src, ref):
        os.remove(f)


if __name__ == "__main__":
    main()
